from pptx import Presentation
from pptx.util import Inches, Pt
import os

prs = Presentation()

# Slide 1: Project Overview
title_slide_layout = prs.slide_layouts[0]
slide1 = prs.slides.add_slide(title_slide_layout)
title = slide1.shapes.title
subtitle = slide1.placeholders[1]
title.text = "BTP Mid-Term Report"
subtitle.text = "Multi-Robot Task Orchestration for Construction Assembly\nProject Overview"

# Slide 2: Progress Till Now
bullet_slide_layout = prs.slide_layouts[1]
slide2 = prs.slides.add_slide(bullet_slide_layout)
shapes2 = slide2.shapes
title_shape2 = shapes2.title
body_shape2 = shapes2.placeholders[1]
title_shape2.text = "Progress Till Now"
tf = body_shape2.text_frame
tf.text = "Simulation Environment: Fully functional PyBullet setup"
p = tf.add_paragraph()
p.text = "Robot Control & Kinematics: 7-DOF Panda arm, grasp-error compensation"
p = tf.add_paragraph()
p.text = "Vision Integration: 3D CV pipeline for overhead camera detection"
p = tf.add_paragraph()
p.text = "Task Orchestration & LLM Brain: Not started yet"

# Add image to Slide 2
img_path_scene = r"C:\Users\Nawang Tashi\.gemini\antigravity-ide\brain\cbbac865-4056-4654-a0c0-6439bfbecf31\scene_placement_1790571874278.jpg"
if os.path.exists(img_path_scene):
    slide2.shapes.add_picture(img_path_scene, Inches(5.5), Inches(2), width=Inches(4))

# Slide 3: Current Challenges & Problems
slide3 = prs.slides.add_slide(bullet_slide_layout)
shapes3 = slide3.shapes
title_shape3 = shapes3.title
body_shape3 = shapes3.placeholders[1]
title_shape3.text = "Current Challenges"
tf3 = body_shape3.text_frame
tf3.text = "Workspace Constraints"
p3 = tf3.add_paragraph()
p3.text = "Primary robotic arm operates in a fixed, static workspace"
p3 = tf3.add_paragraph()
p3.text = "Limits maximum size and height of the wall"
p3 = tf3.add_paragraph()
p3.text = "Transitioning to dynamic base placement or mobile base is complex"

# Add image to Slide 3
img_path_vision = r"C:\Users\Nawang Tashi\.gemini\antigravity-ide\brain\cbbac865-4056-4654-a0c0-6439bfbecf31\vision_planning_1790571889723.jpg"
if os.path.exists(img_path_vision):
    slide3.shapes.add_picture(img_path_vision, Inches(5.5), Inches(2), width=Inches(4))

# Slide 4: Future Goals
slide4 = prs.slides.add_slide(bullet_slide_layout)
shapes4 = slide4.shapes
title_shape4 = shapes4.title
body_shape4 = shapes4.placeholders[1]
title_shape4.text = "Future Goals (Next 2 Months)"
tf4 = body_shape4.text_frame
tf4.text = "Refine LLM Orchestration: Transition to prompt-engineered LLM brain"
p4 = tf4.add_paragraph()
p4.text = "Improve Physics Realism: Realistic kinematics and dynamic grasping"
p4 = tf4.add_paragraph()
p4.text = "Comprehensive Benchmarking: Automated evaluations and Gantt charts"
p4 = tf4.add_paragraph()
p4.text = "Final Project Deliverables: Documentation, defense, and code release"

prs.save("BTP_MidTerm_Presentation.pptx")
print("Presentation created successfully: BTP_MidTerm_Presentation.pptx")
